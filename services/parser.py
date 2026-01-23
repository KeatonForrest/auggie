"""parser.py - Extract text from uploaded documents."""

import io
from pathlib import Path

# PDF
from pypdf import PdfReader

# DOCX
from docx import Document as DocxDocument

# PPTX
from pptx import Presentation


class DocumentParser:
    """Service for extracting text from various document formats."""

    SUPPORTED_TYPES = {'pdf', 'docx', 'pptx', 'txt', 'md'}

    def parse(self, file_bytes: bytes, file_type: str) -> str:
        """Extract text from a document.

        Args:
            file_bytes: Raw file content
            file_type: File extension (pdf, docx, pptx, txt, md)

        Returns:
            Extracted text content

        Raises:
            ValueError: If file type is not supported
        """
        file_type = file_type.lower().lstrip('.')

        if file_type not in self.SUPPORTED_TYPES:
            raise ValueError(
                f"Unsupported file type: {file_type}. "
                f"Supported: {', '.join(sorted(self.SUPPORTED_TYPES))}"
            )

        if file_type == 'pdf':
            return self._parse_pdf(file_bytes)
        elif file_type == 'docx':
            return self._parse_docx(file_bytes)
        elif file_type == 'pptx':
            return self._parse_pptx(file_bytes)
        elif file_type in ('txt', 'md'):
            return file_bytes.decode('utf-8', errors='ignore')

        return ""

    def _parse_pdf(self, file_bytes: bytes) -> str:
        """Extract text from PDF."""
        reader = PdfReader(io.BytesIO(file_bytes))
        text_parts = []

        for page_num, page in enumerate(reader.pages, 1):
            text = page.extract_text()
            if text and text.strip():
                text_parts.append(f"[Page {page_num}]\n{text}")

        return "\n\n".join(text_parts)

    def _parse_docx(self, file_bytes: bytes) -> str:
        """Extract text from DOCX."""
        doc = DocxDocument(io.BytesIO(file_bytes))
        text_parts = []

        # Extract paragraphs
        for para in doc.paragraphs:
            if para.text.strip():
                # Check if it's a heading
                if para.style and para.style.name.startswith('Heading'):
                    text_parts.append(f"\n## {para.text}\n")
                else:
                    text_parts.append(para.text)

        # Extract from tables
        for table in doc.tables:
            table_text = []
            for row in table.rows:
                row_text = ' | '.join(
                    cell.text.strip() for cell in row.cells if cell.text.strip()
                )
                if row_text:
                    table_text.append(row_text)
            if table_text:
                text_parts.append("\n" + "\n".join(table_text) + "\n")

        return "\n".join(text_parts)

    def _parse_pptx(self, file_bytes: bytes) -> str:
        """Extract text from PowerPoint."""
        prs = Presentation(io.BytesIO(file_bytes))
        text_parts = []

        for slide_num, slide in enumerate(prs.slides, 1):
            slide_texts = [f"[Slide {slide_num}]"]

            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    slide_texts.append(shape.text)

                # Extract from tables in slides
                if shape.has_table:
                    for row in shape.table.rows:
                        row_text = ' | '.join(
                            cell.text.strip() for cell in row.cells if cell.text.strip()
                        )
                        if row_text:
                            slide_texts.append(row_text)

            if len(slide_texts) > 1:  # Has content beyond slide number
                text_parts.append("\n".join(slide_texts))

        return "\n\n".join(text_parts)


def get_file_type(filename: str) -> str:
    """Extract file type from filename.

    Args:
        filename: Original filename

    Returns:
        File extension without the dot (e.g., 'pdf', 'docx')
    """
    return Path(filename).suffix.lower().lstrip('.')
